import asyncio
import glob
import os
from docx import Document
from dotenv import load_dotenv
from langchain_google_genai import ChatGoogleGenerativeAI

load_dotenv()

# =========================================================================
# 輔助函式：從不同的檔案格式中抽出「純文字」
# =========================================================================
def extract_text_from_file(file_path):
    ext = os.path.splitext(file_path)[1].lower()

    # 1. Word 檔案 (.docx)
    if ext == ".docx":
        try:
            doc = Document(file_path)
            full_text = []
            for para in doc.paragraphs:
                full_text.append(para.text)
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        full_text.append(cell.text)
            return "\n".join(full_text)
        except Exception as e:
            print(f"❌ 讀取 Word 檔失敗 {file_path}: {e}")
            return ""

    # 2. PDF 檔案 (.pdf)
    elif ext == ".pdf":
        try:
            import pypdf
            reader = pypdf.PdfReader(file_path)
            text_list = []
            for page in reader.pages:
                extracted = page.extract_text()
                if extracted:
                    text_list.append(extracted)
            return "\n".join(text_list)
        except Exception as e:
            print(f"❌ 讀取 PDF 檔失敗 {file_path}: {e}")
            return ""

    # 3. Markdown 或普通文字檔 (.md, .txt)
    elif ext in [".md", ".txt"]:
        try:
            with open(file_path, "r", encoding="utf-8") as f:
                return f.read()
        except Exception as e:
            print(f"❌ 讀取文字檔失敗 {file_path}: {e}")
            return ""

    return ""

# =========================================================================
# 輔助函式：自動檢查既存檔案並生成不重複的流水號檔名
# =========================================================================
def get_unique_filepath(directory, prefix, extension=".md"):
    """
    自動檢查資料夾內既存檔案，回傳下一個不重複的流水號路徑。
    範例：若已存在 B1_Spec_Review_v1.md，則會自動產出 B1_Spec_Review_v2.md
    """
    i = 1
    while True:
        filename = f"{prefix}_v{i}{extension}"
        filepath = os.path.join(directory, filename)
        if not os.path.exists(filepath):
            return filepath
        i += 1

# =========================================================================
# 核心邏輯：資深 QA 視角 - 提煉【功能說明】、【B1 規格審查】與【測試重點】
# =========================================================================
async def summarize_spec_via_ai(file_name, raw_text):
    # llm = ChatGoogleGenerativeAI(model="gemini-3.6-flash", temperature=0.1)
    llm = ChatGoogleGenerativeAI(model="gemini-2.5-flash", temperature=0.1)

    prompt = f"""
你是一位極度嚴謹且具備 10 年以上軟體測試經驗的資深 QA 架構師。
請用專業 QA 的視角，深度解構以下這份原始 Spec / 需求內容。你的任務不是單純摘要，而是要發掘規格中的盲點與潛在測試場景，並將其轉化為結構化的測試綱要。

【檔案名稱】: {file_name}
【原始文件內容】:
{raw_text}

----------------------------------------------------------------
【QA 深度精煉規則】:
1. 你必須為這份文件產出三個核心區塊：『【功能說明】』、『【規格審查與漏洞分析 (B1 Spec Review)】』與『【測試重點】』。
2. 『【功能說明】』：請用簡潔易懂的文字，說明這個模組/功能的核心目標與商業邏輯。
3. 『【規格審查與漏洞分析 (B1 Spec Review)】』：請發揮 QA 質質疑精神，列出此規格本身模糊、矛盾、未定義或具資安風險處（如：極限邊界未規範、UI狀態未重置、輸入驗證缺漏、XSS/SQLi防護未提及等）。
4. 『【測試重點】』：運用 QA 思維（等價類劃分、邊界值分析、異常防呆、狀態轉換）列出具體測試方向。包含：
    - [正向] (Happy Path) 核心流程與預期結果
    - [異常] (Negative/Error Handling) 操作錯誤、必填漏填、格式錯誤等
    - [邊界] (Edge Cases) 負數、小數點、字數極限、特殊字元等
    - [狀態/權限] (State/Permission) 票種/狀態切換、資料殘留與權限控制
5. 測試重點請用條列式呈現，每一條都要具備「驗證...是否...」的具體行動與預期結果。
6. 請直接輸出結果，不要包含額外的問候語或 Markdown 程式碼包裹標籤（如 ```markdown）。

請嚴格輸出如下格式：
# [功能模組名稱] - {file_name}

【功能說明】
(在此填寫提煉後的功能說明)

【規格審查與漏洞分析 (B1 Spec Review)】
1. [規格模糊] ...
2. [邊界未定義] ...
3. [資安/防呆缺漏] ...

【測試重點】
1. [正向] 驗證...
2. [異常] 驗證...
3. [邊界] 驗證...
4. [狀態/權限] 驗證...
"""

    try:
        response = await llm.ainvoke(prompt)
        content = response.content

        if isinstance(content, list):
            text_parts = []
            for block in content:
                if isinstance(block, str):
                    text_parts.append(block)
                elif isinstance(block, dict) and "text" in block:
                    text_parts.append(block["text"])
            content = "\n".join(text_parts)

        content = content.strip()
        if content.startswith("```markdown"):
            content = content[11:]
        if content.startswith("```"):
            content = content[3:]
        if content.endswith("```"):
            content = content[:-3]

        return content.strip()
    except Exception as e:
        print(f"❌ AI 分析失敗: {e}")
        return ""

# =========================================================================
# 主流程入口
# =========================================================================
async def main():
    BASE_DIR = os.path.dirname(os.path.abspath(__file__))  # ai_agents/ 資料夾
    ROOT_DIR = os.path.dirname(BASE_DIR)                   # 專案根目錄

    specs_dir = os.path.join(BASE_DIR, "pending_specs")
    docs_dir = os.path.join(ROOT_DIR, "docs")
    os.makedirs(docs_dir, exist_ok=True)

    # 固定寫入 feature.md 供後續腳本讀取
    feature_file = os.path.join(specs_dir, "feature.md")
    
    # 動態取得不重複的流水號路徑 (如 B1_Spec_Review_v1.md)
    b1_review_file = get_unique_filepath(docs_dir, "B1_Spec_Review")

    # 1. 掃描 pending_specs 資料夾底下的所有原始規格檔案
    all_files = []
    for ext in ["*.docx", "*.pdf", "*.md", "*.txt"]:
        all_files.extend(glob.glob(os.path.join(specs_dir, ext)))

    # 排除 feature.md 以及被標記為 .processed 的檔案
    pending_files = [f for f in all_files if "feature.md" not in f and not f.endswith(".processed")]

    if not pending_files:
        print(f"⚠️ 在 {specs_dir}/ 中沒有找到任何需要提煉的原始 Spec 檔案。")
        return

    print(f"📦 偵測到 {len(pending_files)} 個原始文件，開始一個個餵給 AI 進行 QA 視角精煉...\n")

    combined_features = []

    # 2. 一個個讀取、分析並收集
    for file_path in pending_files:
        file_name = os.path.basename(file_path)
        print(f"📄 正在讀取文件: {file_name} ...")

        raw_text = extract_text_from_file(file_path)
        if not raw_text.strip():
            print(f"⚠️ 檔案 {file_name} 內容為空或讀取失敗，跳過。")
            continue

        print(f"🤖 正在讓 AI 資深 QA 深度分析【功能說明】、【B1 規格審查】與【測試重點】...")
        refined_text = await summarize_spec_via_ai(file_name, raw_text)

        if refined_text:
            combined_features.append(refined_text)
            print(f"✅ {file_name} 精煉成功！")
        print("-" * 50)

    # 3. 雙向寫入 feature.md 與流水號的 B1_Spec_Review_vX.md
    if combined_features:
        final_markdown_content = "\n\n---\n\n".join(combined_features)

        # 寫入 ai_agents/pending_specs/feature.md
        with open(feature_file, "w", encoding="utf-8") as f:
            f.write(final_markdown_content)

        # 寫入 docs/B1_Spec_Review_vX.md
        with open(b1_review_file, "w", encoding="utf-8") as f:
            f.write(final_markdown_content)

        print(f"\n🎉 【全自動精煉完成】所有的 Spec 已成功被 QA 大腦分析完畢！")
        print(f"📝 雙向同步寫入完成：")
        print(f"   - {feature_file}")
        print(f"   - {b1_review_file}")
        print(f"🚀 現在您可以放心地直接執行 `python3 ai_agents/generate_qase_bdd.py` 來批量產生案例並同步至 Qase！")

if __name__ == "__main__":
    asyncio.run(main())